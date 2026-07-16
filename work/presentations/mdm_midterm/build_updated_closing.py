from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
QA = HERE / "scoring_qa"
OUT = ROOT / "主数据管理_中期汇报_更新版.pptx"

W, H = 13.333333, 7.5
C = {
    "bg": "F5FAFC", "ink": "102033", "muted": "5E7188", "blue": "1456A0",
    "cyan": "06A7C7", "teal": "0E8F86", "green": "1F9D55", "amber": "F59E0B",
    "white": "FFFFFF", "line": "D5E6EE", "soft": "E8F4F8",
}


def rgb(value):
    return RGBColor.from_string(value)


def add_text(slide, value, x, y, w, h, size=18, bold=False, color="ink", align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.text = value
    p.alignment = align
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = rgb(C.get(color, color))
    return shape


def add_box(slide, x, y, w, h, fill="white", line="line", radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(C.get(fill, fill))
    shape.line.color.rgb = rgb(C.get(line, line))
    shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def header(slide, section, title, subtitle):
    add_text(slide, section, .67, .36, 5.8, .28, 14, True, "cyan")
    add_text(slide, title, .67, .75, 11.2, .58, 27, True)
    add_box(slide, .67, 1.48, .96, .04, "cyan", "cyan", False)
    add_text(slide, subtitle, .67, 1.62, 11.0, .42, 15, False, "muted")


def footer(slide, page):
    add_text(slide, "基于 Dify 的客户/商品主数据自动化匹配项目", .67, 7.02, 7.0, .24, 10.5, False, "71859A")
    add_text(slide, f"{page:02d}", 11.9, 7.0, .75, .26, 12, True, "blue", PP_ALIGN.RIGHT)


def metric(slide, label, value, note, x, y, accent):
    add_box(slide, x, y, 2.60, 1.38)
    add_box(slide, x, y, .06, 1.38, accent, accent, False)
    add_text(slide, value, x + .23, y + .12, 2.1, .42, 24, True)
    add_text(slide, label, x + .23, y + .62, 2.1, .28, 13, True, "blue")
    add_text(slide, note, x + .23, y + .93, 2.08, .34, 10.5, False, "muted")


def card(slide, title, body, x, y, w, h, accent):
    add_box(slide, x, y, w, h)
    add_text(slide, title, x + .25, y + .15, w - .5, .3, 15.5, True, accent)
    add_text(slide, body, x + .25, y + .52, w - .5, h - .62, 12, False, "muted")


def step(slide, number, title, body, x, color):
    add_box(slide, x, 2.54, 1.98, 1.16)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + .18), Inches(2.74), Inches(.36), Inches(.36))
    circle.fill.solid(); circle.fill.fore_color.rgb = rgb(C[color]); circle.line.color.rgb = rgb(C[color])
    add_text(slide, str(number), x + .18, 2.75, .36, .30, 11, True, "white", PP_ALIGN.CENTER)
    add_text(slide, title, x + .65, 2.66, 1.15, .28, 13, True)
    add_text(slide, body, x + .65, 3.00, 1.15, .52, 9.5, False, "muted")


def new_slide(prs, fill="bg"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(C[fill])
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    for index in range(1, 11):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(QA / f"slide-{index:02d}.png"), 0, 0, width=prs.slide_width, height=prs.slide_height)

    slide = new_slide(prs)
    header(slide, "10  前端量化展示", "从数据接入到主数据结果，关键状态在一个界面内可量化、可追踪", "处理规模、匹配质量、运行状态与结果审计集中展示，让主数据治理从黑盒任务变为可观察流程。")
    metric(slide, "来源记录覆盖", "100%", "A 1,363 条 + B 3,226 条均完成归属", .75, 2.42, "green")
    metric(slide, "候选压缩率", "99.5%", "约 440 万组合压缩为 22,208 个候选", 3.70, 2.42, "cyan")
    metric(slide, "跨表匹配精度", "0.784", "638 条正确链接，结果可回查来源", 6.65, 2.42, "blue")
    metric(slide, "统一实体视图", "4,589", "黄金记录、来源映射与决策日志联动", 9.60, 2.42, "teal")
    card(slide, "运行看板", "展示批次进度、处理规模、成功率与异常状态，快速判断任务是否正常完成。", .82, 4.43, 3.40, 1.40, "blue")
    card(slide, "匹配结果", "支持按实体、来源表和置信度筛选，查看自动合并、待审与未匹配记录。", 4.97, 4.43, 3.40, 1.40, "teal")
    card(slide, "质量追踪", "覆盖率、唯一性、Precision / Recall 与候选规模集中呈现，便于横向比较。", 9.12, 4.43, 3.40, 1.40, "green")
    footer(slide, 11)

    slide = new_slide(prs)
    header(slide, "11  万能匹配", "任意两张业务表，一次完成字段理解、候选生成与实体归并", "无需预先固定表结构；字段映射、机器学习评分和 LLM 冲突判定组成可复用流水线。")
    items = [
        ("任意表输入", "上传 CSV / Excel，读取字段、类型与样例"),
        ("字段语义映射", "识别名称、品牌、价格、地址等同义字段"),
        ("候选召回评分", "Blocking 压缩规模，Baseline 输出置信度"),
        ("冲突簇判定", "LLM 处理一对多、多对一与缺失字段"),
        ("统一结果输出", "生成黄金记录、来源链接与审计决策"),
    ]
    colors = ["blue", "teal", "cyan", "amber", "green"]
    for i, ((title, body), color) in enumerate(zip(items, colors)):
        x = .73 + i * 2.50
        step(slide, i + 1, title, body, x, color)
        if i < 4:
            add_text(slide, "→", x + 2.02, 2.91, .44, .34, 21, True, "muted", PP_ALIGN.CENTER)
    card(slide, "结构通用", "字段名、列顺序不同或部分字段缺失，均可通过 schema 映射转为统一匹配语义。", .82, 4.32, 3.40, 1.43, "blue")
    card(slide, "策略通用", "规则处理高置信样本，机器学习负责排序，LLM 只进入复杂冲突簇。", 4.97, 4.32, 3.40, 1.43, "amber")
    card(slide, "结果通用", "统一输出实体、黄金记录、来源映射和判定日志，可接入客户、商品等数据域。", 9.12, 4.32, 3.40, 1.43, "green")
    add_text(slide, "核心价值：一次配置完成从“任意输入表”到“可信主数据”的端到端匹配闭环。", 1.3, 6.10, 10.7, .34, 17, True, "ink", PP_ALIGN.CENTER)
    footer(slide, 12)

    slide = new_slide(prs)
    header(slide, "12  总结", "已完成可运行、可量化、可扩展的主数据自动匹配原型", "中期成果覆盖算法、工作流、后端服务、前端展示与质量审计，验证了技术路线和工程落地的可行性。")
    card(slide, "已经完成", "Baseline 模型、候选压缩、分层决策、Dify 工作流、FastAPI 服务、主数据表结构与可视化前端已形成端到端闭环。", .82, 2.40, 3.40, 2.02, "blue")
    card(slide, "阶段价值", "4,589 条来源记录实现 100% 覆盖；候选规模压缩至 22,208；结果具备来源映射、质量指标和决策追踪。", 4.97, 2.40, 3.40, 2.02, "teal")
    card(slide, "下一步", "围绕误合并与漏匹配调优阈值和特征，补充多数据集对比实验，并完善人工复核与生产部署能力。", 9.12, 2.40, 3.40, 2.02, "green")
    add_box(slide, 1.34, 5.16, 10.66, .94, "soft", "line")
    add_text(slide, "结论", 1.68, 5.40, 1.1, .3, 17, True, "blue")
    add_text(slide, "项目已从“匹配模型”推进为“主数据治理系统”：既能处理复杂匹配，又能量化效果、解释结果并持续迭代。", 2.95, 5.34, 8.4, .42, 15.5, True)
    footer(slide, 13)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
